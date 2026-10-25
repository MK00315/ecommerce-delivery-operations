"""Read-only saved-file checks after Artifact Tool authoring and native Excel QA."""
from datetime import datetime
import hashlib
import json
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/native"
path = OUT / "Delivery_Operations.xlsx"
values = load_workbook(path, data_only=True)
formulas = load_workbook(path, data_only=False)
assert values.sheetnames == ['Overview','Review capacity','Statistics','Sources','ReadMe']
assert values.active.title == 'Overview'
assert values['Overview']['C7'].value == 96470
assert values['Overview']['C8'].value == 6534
assert abs(values['Overview']['C9'].value - 6534 / 96470) < 1e-12
assert values['Overview']['G7'].value == 99441 and values['Overview']['G8'].value == 2971
assert values['Review capacity']['C6'].value == 'capacity_120_cost_1'
assert [values['Review capacity'].cell(13,c).value for c in range(3,7)] == [721,3314,24,79]
assert values['Review capacity']['K29'].value == 30
assert all(values['ReadMe'].cell(r,5).value == 0 for r in range(30,36))
errors = [(s.title,c.coordinate,c.value) for s in values for row in s for c in row if c.data_type == 'e']
assert not errors
assert sum(len(s._charts) for s in formulas) == 3
assert sum(len(s.tables) for s in formulas) == 5
assert sum(len(s.data_validations.dataValidation) for s in formulas) == 5
assert not any(s.sheet_view.showGridLines for s in formulas)
chart_refs = []
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None
    assert not any(name.startswith('xl/externalLinks/') for name in archive.namelist())
    for name in archive.namelist():
        if name.startswith('xl/charts/chart') and name.endswith('.xml'):
            document = ET.fromstring(archive.read(name))
            chart_refs.append({'chart':name,'formulas':[e.text for e in document.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}f')]})
            assert document.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}f')
    line = ET.fromstring(archive.read('xl/charts/chart1.xml'))
    assert line.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').attrib['val'] == '0EA5E9'
verification = {'checked_at':datetime.now().astimezone().isoformat(),'passed':True,
    'check_scope':'Read-only verification of the published saved workbook; no native Excel UI was rerun',
    'historical_native_evidence':'docs/REPORTING_VERIFICATION.json records the original native authoring and scenario checks',
    'worksheets':5,'native_charts':3,'native_tables':5,'data_validation_rules':5,
    'source_orders_represented':99441,'eligible_orders':96470,'late_orders':6534,
    'saved_formula_errors':errors,'reconciliation_checks_passed':6,
    'chart_bindings':chart_refs,'source_data_sampled':False,
    'snapshot_limit':'Scenario selector reads executed Python runs. Uniform staffing calculator recalculates; no live optimizer or cloud refresh in the workbook.',
    'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size}
(OUT / 'EXCEL_VERIFICATION.json').write_text(json.dumps(verification,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:verification[k] for k in ['passed','worksheets','native_charts','native_tables','source_orders_represented','bytes']},indent=2))
